#!/usr/bin/env python3
"""Build a fillable Word document from task VII (the calculation tasks) of the exam sheet.

Abel's request (2026-10-02): "ebbol csinalj a szamitasos feladatokbol egy kitoltheto word-ot".

WHY THE TEXT IS NOT RETYPED: this is an exam sheet, and every figure in it decides the
answer. The task text is sliced out of the PDF's own extraction (pdfplumber), never typed
by hand, so a digit cannot drift between the source and the output.

AND THE GUARANTEE IS NOT THAT CLAIM, BUT THE CHECK AT THE END: the script extracts every
number token from pages 10-14 of the PDF and from the generated .docx, and compares the two
multisets. If a single figure is missing or extra, it says so and names it. A document that
LOOKS like the exam but carries one wrong digit is worse than no document.

WHAT "FILLABLE" MEANS HERE: python-docx cannot emit native Word form fields, so each
sub-task gets a single-cell bordered table underneath it. In Word that is a click-and-type
box, it survives editing, and it prints as a visible answer field. Items 8 and 12 ask for
several values, so they get one box per value, labelled.

Usage: python3 scripts/vizsga-vii-kitoltheto-docx.py
"""
import io
import re
import sys

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

PDF = "store/vizsga-feladatsor/Kereskedo_es_webaruhazi_technikus_Feladatsor_2026.pdf"
KI = "store/vizsga-feladatsor/VII_szamitasi_feladatok_2026_kitoltheto.docx"
ELSO, UTOLSO = 10, 14          # a VII. feladat oldalai, 1-tol szamozva

FEJLEC_MINTA = re.compile(
    r"^(2019\. évi LXXX\. törvény.*|37\s*5 0416.*|Versenyzői kód:.*|5 0416 13 03.*|\d+/\d+)$")


def vii_szovege():
    """A VII. feladat sorai, a PDF sajat kivonatabol, fejlec- es labsorok nelkul."""
    sorok = []
    with pdfplumber.open(PDF) as pdf:
        for i in range(ELSO - 1, UTOLSO):
            for sor in (pdf.pages[i].extract_text() or "").split("\n"):
                s = sor.strip()
                if not s or FEJLEC_MINTA.match(s):
                    continue
                sorok.append(s)
    return sorok


def szamok(szoveg):
    """Minden szam-token, a magyar ezres szokozzel es tizedesvesszovel egyutt."""
    t = re.sub(r"(?<=\d)[  ](?=\d)", "", szoveg)        # "1 470" -> "1470"
    return sorted(re.findall(r"\d+(?:[.,]\d+)?", t))


def valaszmezo(doc, cimke=None):
    if cimke:
        p = doc.add_paragraph()
        r = p.add_run(cimke)
        r.bold = True
        r.font.size = Pt(9)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    c.text = ""
    c.paragraphs[0].add_run(" ").font.size = Pt(14)
    doc.add_paragraph()


def main():
    sorok = vii_szovege()
    forras = "\n".join(sorok)

    doc = Document()
    cim = doc.add_heading("VII. Feladat (15 pont)", level=1)
    cim.alignment = WD_ALIGN_PARAGRAPH.LEFT
    al = doc.add_paragraph()
    r = al.add_run("Numerikus értékek megadása (Számítási feladatok)")
    r.bold = True

    e = doc.add_paragraph()
    r = e.add_run(
        "Forrás: Magyar Kereskedelmi és Iparkamara, Országos Szakmai Tanulmányi Verseny, "
        "területi előválogató, központi interaktív feladatsor, 2026. "
        "Szakma: 5 0416 13 03 Kereskedő és webáruházi technikus. "
        "A feladatok szövege a hivatalos PDF 10-14. oldaláról származik, változtatás nélkül."
    )
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    doc.add_paragraph()

    # A sorokat a "<n>. <cim>" alaku fejlecek mentén vágjuk feladatokra.
    uj_feladat = re.compile(r"^(\d{1,2})\.\s+([A-ZÁÉÍÓÖŐÚÜŰ].*)$")
    blokkok, aktualis = [], None
    for s in sorok:
        m = uj_feladat.match(s)
        if m and not s.lower().startswith(("a)", "b)", "c)")):
            if aktualis:
                blokkok.append(aktualis)
            aktualis = {"szam": m.group(1), "cim": m.group(2), "sorok": []}
        elif aktualis:
            aktualis["sorok"].append(s)
    if aktualis:
        blokkok.append(aktualis)

    # A VII. fejlécsorai a legelso blokk ELOTT allnak, azokat mar kiirtuk.
    for b in blokkok:
        h = doc.add_heading(f"{b['szam']}. {b['cim']}", level=2)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for s in b["sorok"]:
            if re.match(r"^\.{0,4}…{0,2}\.{0,4}\s*(pont)?\s*/\s*\d+\s*pont$", s) or \
               re.match(r"^…+\s*/\s*\d+\s*pont$", s):
                p = doc.add_paragraph()
                r = p.add_run(s)
                r.italic = True
                r.font.size = Pt(8)
                continue
            doc.add_paragraph(s)
        # Tobb-valaszos feladatok: 8. (harom ertek) es 12. (ket ertek)
        if b["szam"] == "8":
            for cimke in ("Értékesítés nettó árbevétele:", "Anyag jellegű ráfordítás:",
                          "Személyi jellegű ráfordítás:"):
                valaszmezo(doc, cimke)
        elif b["szam"] == "12":
            for cimke in ("Munkavállaló közterhe:", "Munkáltató közterhe:"):
                valaszmezo(doc, cimke)
        else:
            valaszmezo(doc, "Válasz:")

    doc.save(KI)

    # --- A GARANCIA: minden szam-token megvan-e, es nem keletkezett-e tobblet ---
    ujra = Document(KI)
    kimenet = "\n".join(p.text for p in ujra.paragraphs)
    for t in ujra.tables:
        for sor in t.rows:
            for c in sor.cells:
                kimenet += "\n" + c.text

    a, b_ = szamok(forras), szamok(kimenet)
    from collections import Counter
    hianyzo = Counter(a) - Counter(b_)
    tobblet = Counter(b_) - Counter(a)
    print(f"kiirva: {KI}")
    print(f"feladat-blokk: {len(blokkok)}")
    print(f"szam-token a forrasban: {len(a)} | a .docx-ben: {len(b_)}")
    if hianyzo:
        print("HIANYZO SZAMOK:", dict(hianyzo)); sys.exit(1)
    if tobblet:
        print("TOBBLET SZAMOK (a sablon sajat szamai is ide esnek):", dict(tobblet))
    print("MINDEN FORRAS-SZAM MEGVAN a kimenetben.")


if __name__ == "__main__":
    main()
