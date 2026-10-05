#!/usr/bin/env python3
"""Anomaly scan for the K&F filter pricing workbook David returns.

David's request (2026-10-02, Telegram 8218): find mispricings -- order-of-magnitude
typos (29990 vs 2990), outliers, and families where the price stops following the
pattern as the filter diameter grows.

Only rows flagged 'i' in the "arazni" column are in scope; 'a' means the item will be
cleared with a promotion instead and David said to leave those alone.

The model, read from the sheet's own formulas (NOT assumed):
    G = purchase price, USD
    T = UJ PDB ARRES      = W / (G * 330)      <- note: 330, not the Parameterek sheet's rate
    U = UJ KISKER BRUTTO  <- HAND-TYPED in most rows, so this is where a typo lands
    V = UJ KISKER NETTO   = U / 1.27
    W = UJ NAGYKER        = V - V*Y
    X = UJ mikrosat atadoi= V * 0.74
    Z = mikrosat PDB arres= X / (G * 330)

Usage: python3 scripts/kf-arazas-anomalia.py <xlsx>
"""
import collections
import re
import statistics
import sys

import openpyxl

ARFOLYAM_KEPLETBEN = 330.0


def szam(v):
    return v if isinstance(v, (int, float)) else None


def meret_mm(nev):
    """Filter diameter in mm from the product name, or None.

    Only a standalone 2-3 digit mm value counts. A plate filter ("100*150*2 mm")
    has no single diameter, so those names are deliberately left unparsed.
    """
    if not nev:
        return None
    if re.search(r"\d+\s*\*\s*\d+", nev):
        return None
    # A tizedes alakot MINDKET irasmoddal kell fogni: "40.5mm" es "40,5mm" is elofordul.
    # A vesszos alak kihagyasa nemán kiejtene a sort a monotonitas-vizsgalatbol.
    talalat = re.findall(r"(?<![\d.,*/])(\d{2,3}(?:[.,]\d)?)\s*mm\b", nev, re.IGNORECASE)
    ertekek = {float(t.replace(",", ".")) for t in talalat}
    if len(ertekek) != 1:
        return None
    return ertekek.pop()


def csaladnev(nev):
    """The product name with the diameter removed, so one size-series groups together."""
    if not nev:
        return None
    s = re.sub(r"(?<![\d.,*/])\d{2,3}(?:[.,]\d)?\s*mm\b", "<MM>", nev, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", s).strip().lower()


def main():
    utvonal = sys.argv[1] if len(sys.argv) > 1 else (
        "store/kf-arazas/KF_szuroarazas_2026-09-22_rendezett_DAVIDTOL_2026-10-02.xlsx")

    wb_k = openpyxl.load_workbook(utvonal, data_only=False)
    wb_v = openpyxl.load_workbook(utvonal, data_only=True)
    ws_k, ws_v = wb_k["Sheet1"], wb_v["Sheet1"]
    fejlec = [c.value for c in ws_v[1]]
    oszlop = {h: i for i, h in enumerate(fejlec)}

    C = dict(
        cikkszam=oszlop["Cikkszám"], szall=oszlop["Szállítói cikkszám"],
        nev=oszlop["Megnevezés"], arazni=oszlop["árazni"],
        usd=oszlop["árlista beszer 2026.09.30"],
        regi_brutto=oszlop["régi bruttó kisker"],
        fotoplus=oszlop["Fotoplus kisker bruttó"], kamerapro=oszlop["Kamerapro kisker bruttó"],
        arres=oszlop["ÚJ PDB ÁRRÉS"], brutto=oszlop["ÚJ KISKER BRUTTÓ"],
        netto=oszlop["ÚJ KISKER NETTÓ"], nagyker=oszlop["ÚJ NAGYKER"],
        mikrosat=oszlop[" ÚJ mikrosat átadói"], kisker_arres=oszlop["ÚJ KISKER ÁRRÉS"],
        mikrosat_arres=oszlop["mikrosat PDB árrés"], kifuto=oszlop["Kifutó"],
    )

    sorok = []
    for ertek, keplet in zip(ws_v.iter_rows(min_row=2), ws_k.iter_rows(min_row=2)):
        v = [c.value for c in ertek]
        if not any(x not in (None, "") for x in v):
            continue
        if v[C["arazni"]] != "i":
            continue
        sorok.append({
            "excel_sor": ertek[0].row,
            "cikkszam": v[C["cikkszam"]], "szall": v[C["szall"]], "nev": v[C["nev"]],
            "usd": szam(v[C["usd"]]), "regi_brutto": szam(v[C["regi_brutto"]]),
            "fotoplus": szam(v[C["fotoplus"]]), "kamerapro": szam(v[C["kamerapro"]]),
            "arres": szam(v[C["arres"]]), "brutto": szam(v[C["brutto"]]),
            "netto": szam(v[C["netto"]]), "nagyker": szam(v[C["nagyker"]]),
            "mikrosat": szam(v[C["mikrosat"]]),
            "kisker_arres": szam(v[C["kisker_arres"]]),
            "mikrosat_arres": szam(v[C["mikrosat_arres"]]),
            "kifuto": v[C["kifuto"]],
            "brutto_keplet": keplet[C["brutto"]].value if isinstance(
                keplet[C["brutto"]].value, str) and str(
                keplet[C["brutto"]].value).startswith("=") else None,
        })

    print(f"HATOKOR: {len(sorok)} sor (arazni='i')\n")
    lelet = collections.OrderedDict()

    # 1. Belso konzisztencia: a keplet-oszlopok ertekei egyeznek-e a sajat kepletukkel.
    baj = []
    for r in sorok:
        if None in (r["brutto"], r["netto"], r["nagyker"], r["mikrosat"], r["kisker_arres"]):
            baj.append((r, "hianyzo ertek"))
            continue
        if abs(r["netto"] - r["brutto"] / 1.27) > 1:
            baj.append((r, f"netto={r['netto']:.0f} != brutto/1.27={r['brutto']/1.27:.0f}"))
        elvart_nagyker = r["netto"] - r["netto"] * r["kisker_arres"]
        if abs(r["nagyker"] - elvart_nagyker) > 1:
            baj.append((r, f"nagyker={r['nagyker']:.0f} != netto*(1-arres)={elvart_nagyker:.0f}"))
        if abs(r["mikrosat"] - r["netto"] * 0.74) > 1:
            baj.append((r, f"mikrosat={r['mikrosat']:.0f} != netto*0.74={r['netto']*0.74:.0f}"))
    lelet["1. BELSO KONZISZTENCIA (keplet-oszlopok)"] = baj

    # 2. Nagysagrendi elutes: a PDB arres a kezzel beirt brutto egyetlen fuggvenye.
    arresek = [r["arres"] for r in sorok if r["arres"]]
    med = statistics.median(arresek)
    szoras = statistics.pstdev(arresek)
    baj = []
    for r in sorok:
        if r["arres"] is None:
            continue
        if r["arres"] < med / 3 or r["arres"] > med * 3:
            baj.append((r, f"arres={r['arres']:.2f} (median={med:.2f}), brutto={r['brutto']}"))
    lelet[f"2. ARRES-KIUGRAS (median {med:.2f}, szoras {szoras:.2f}; 3x savon kivul)"] = baj

    # 3. Regi/uj brutto nagysagrendi ugras.
    baj = []
    for r in sorok:
        if not (r["regi_brutto"] and r["brutto"]):
            continue
        arany = r["brutto"] / r["regi_brutto"]
        if arany < 0.5 or arany > 2.0:
            baj.append((r, f"regi={r['regi_brutto']:.0f} -> uj={r['brutto']:.0f} ({arany:.2f}x)"))
    lelet["3. REGI->UJ BRUTTO UGRAS (0.5x alatt vagy 2x felett)"] = baj

    # 4. Piaci ar: a mi brutto arunk a Fotoplus/Kamerapro arhoz kepest.
    baj = []
    for r in sorok:
        for bolt in ("fotoplus", "kamerapro"):
            if not (r[bolt] and r["brutto"]):
                continue
            arany = r["brutto"] / r[bolt]
            if arany < 0.6 or arany > 1.6:
                baj.append((r, f"{bolt}={r[bolt]:.0f} vs mi={r['brutto']:.0f} ({arany:.2f}x)"))
    lelet["4. PIACI ARTOL ELTERES (0.6x alatt vagy 1.6x felett)"] = baj

    # 5. Meret-sorozat: nagyobb atmero nem lehet dragabb aron kisebb.
    csaladok = collections.defaultdict(list)
    for r in sorok:
        mm, cs = meret_mm(r["nev"]), csaladnev(r["nev"])
        if mm and cs:
            csaladok[cs].append((mm, r))
    baj = []
    for cs, tetelek in sorted(csaladok.items()):
        if len(tetelek) < 3:
            continue
        tetelek.sort(key=lambda t: t[0])
        for (mm1, r1), (mm2, r2) in zip(tetelek, tetelek[1:]):
            if None in (r1["brutto"], r2["brutto"]):
                continue
            if r2["brutto"] < r1["brutto"]:
                baj.append((r2, f"{mm1}mm={r1['brutto']:.0f} DE {mm2}mm={r2['brutto']:.0f} "
                                f"(csokken), csalad: {cs[:60]}"))
    lelet["5. MERET-SOROZAT MEGTORIK (nagyobb atmero olcsobb)"] = baj

    # 6. Ar-vegzodes: a kezzel beirt ar tipikusan 990-re vegzodik.
    vegek = collections.Counter(int(r["brutto"]) % 1000 for r in sorok if r["brutto"])
    gyakori = {v for v, db in vegek.items() if db >= 5}
    baj = [(r, f"brutto={r['brutto']:.0f} (vegzodes {int(r['brutto'])%1000}, ritka)")
           for r in sorok if r["brutto"] and int(r["brutto"]) % 1000 not in gyakori]
    lelet[f"6. RITKA AR-VEGZODES (gyakoriak: {sorted(gyakori)})"] = baj

    # 7. Azonos megnevezes, elteroo ar.
    nevek = collections.defaultdict(list)
    for r in sorok:
        if r["nev"]:
            nevek[re.sub(r"\s+", " ", r["nev"]).strip().lower()].append(r)
    baj = []
    for nev, tetelek in nevek.items():
        arak = {r["brutto"] for r in tetelek if r["brutto"]}
        if len(tetelek) > 1 and len(arak) > 1:
            baj.append((tetelek[0], f"{len(tetelek)} sor ugyanezzel a nevvel, arak: "
                                    f"{sorted(int(a) for a in arak)}"))
    lelet["7. AZONOS MEGNEVEZES, ELTERO AR"] = baj

    for cim, baj in lelet.items():
        print(f"=== {cim} -> {len(baj)} talalat")
        for r, miert in baj[:40]:
            print(f"   sor {r['excel_sor']:4d} {str(r['szall'] or ''):12s} "
                  f"{str(r['nev'])[:58]:58s} {miert}")
        if len(baj) > 40:
            print(f"   ... es meg {len(baj)-40}")
        print()

    # Kiegeszito mereszek, amik nem sor-szintu leletek.
    print("=== KIEGESZITO MERESEK")
    print(f"   arfolyam a kepletekben: {ARFOLYAM_KEPLETBEN:.0f}")
    for sor in wb_v["Paraméterek"].iter_rows(min_row=2, values_only=True):
        if sor[0] and "USD" in str(sor[0]):
            print(f"   arfolyam a Parameterek lapon: {sor[1]}  <- a kepletek NEM ezt hasznaljak")
    kisker = collections.Counter(r["kisker_arres"] for r in sorok)
    print(f"   UJ KISKER ARRES kulonbozo ertekei: {dict(kisker)}")
    kez = sum(1 for r in sorok if r["brutto_keplet"] is None)
    print(f"   UJ KISKER BRUTTO: {kez} kezzel beirt, {len(sorok)-kez} kepletes")
    for r in sorok:
        if r["brutto_keplet"]:
            print(f"      sor {r['excel_sor']}: {r['brutto_keplet']}")
    nincs_piaci = sum(1 for r in sorok if not r["fotoplus"] and not r["kamerapro"])
    print(f"   piaci ar NELKUL arazott sor: {nincs_piaci} / {len(sorok)}")


if __name__ == "__main__":
    main()
